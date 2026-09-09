from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import create_engine

from app.adapters.database.models import Base
from app.adapters.database.repositories import SqlAlchemyFormRepository
from app.adapters.export.xlsx import XlsxExporter
from app.adapters.storage.local import LocalEvidenceStorage
from app.application.export_forms import ExportForms
from app.application.import_forms import ImportForms
from app.application.query_forms import FormFilters, QueryForms
from app.application.review_forms import ReviewForms
from app.domain.models import ExportStatus, FormField


def setup_confirmed_form(tmp_path: Path):  # type: ignore[no-untyped-def]
    engine = create_engine(f"sqlite:///{tmp_path / 'demo.db'}")
    Base.metadata.create_all(engine)
    repository = SqlAlchemyFormRepository(engine)
    imports = ImportForms(
        repository, repository, repository, LocalEvidenceStorage(tmp_path / "evidence")
    )
    reviews = ReviewForms(
        repository,
        repository,
        authority_read=repository,
        admission=repository,
        known_producers=frozenset({("human-reviewer", "manual-entry-v1")}),
        known_selection_artifacts=frozenset({"recognition-threshold-v1"}),
    )
    image = tmp_path / "scan.png"
    image.write_bytes(b"image")
    evidence = imports.import_image(image, "FORM-0001", "T1", "1", "operator")
    values = {"employee_id": "E001", "total_quantity": 10, "qualified_quantity": 9}
    for index, field_name in enumerate(values):
        repository.add_form_field(FormField(f"FIELD-{index}", "FORM-0001", field_name, {}))
    reviews.confirm(
        "FORM-0001",
        0,
        values,
        "reviewer",
        "confirmed",
        certificate_ids_by_field={field: None for field in values},
        manual_evidence_ids_by_field={field: evidence.file_id for field in values},
    )
    service = ExportForms(repository, XlsxExporter(), QueryForms(repository))
    return service, reviews, repository, evidence.file_id


def test_xlsx_has_four_sheets_persisted_batch_and_reverse_trace(tmp_path: Path) -> None:
    service, _, repository, _ = setup_confirmed_form(tmp_path)

    batch = service.export("OUTPUT", FormFilters(), tmp_path / "exports", "finance")

    workbook = load_workbook(batch.file_path, read_only=True)
    assert workbook.sheetnames == ["正式数据", "异常与复核", "汇总", "导出说明"]
    official_rows = list(workbook["正式数据"].iter_rows(values_only=True))
    headers, data = official_rows
    row = dict(zip(headers, data, strict=True))
    assert row["export_batch_id"] == batch.export_batch_id
    assert row["form_id"] == "FORM-0001"
    assert row["record_version"] == 1
    assert Path(batch.file_path).exists()
    assert len(batch.file_sha256) == 64
    assert repository.list_export_batches()[0].included_records == (("FORM-0001", 1),)
    assert repository.get_form("FORM-0001").export_status is ExportStatus.EXPORTED  # type: ignore[union-attr]


def test_correction_after_export_requires_new_export_and_old_file_remains(tmp_path: Path) -> None:
    service, reviews, repository, evidence_id = setup_confirmed_form(tmp_path)
    first = service.export("OUTPUT", FormFilters(), tmp_path / "exports", "finance")

    corrected = {"employee_id": "E001", "total_quantity": 11, "qualified_quantity": 10}
    reviews.confirm(
        "FORM-0001",
        1,
        corrected,
        "reviewer",
        "correction",
        certificate_ids_by_field={field: None for field in corrected},
        manual_evidence_ids_by_field={field: evidence_id for field in corrected},
    )

    assert repository.get_form("FORM-0001").export_status is ExportStatus.REEXPORT_REQUIRED  # type: ignore[union-attr]
    assert Path(first.file_path).exists()
    second = service.export("OUTPUT", FormFilters(), tmp_path / "exports", "finance")
    assert second.file_path != first.file_path
    assert Path(first.file_path).exists()
    assert Path(second.file_path).exists()
