"""Build editable submission letters from the reviewed plain-text sources."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'submission'


def new_document(title):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.85)
    section.left_margin = section.right_margin = Inches(0.9)
    for name in ('Normal', 'Title', 'List Bullet'):
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        for attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
            style.element.get_or_add_rPr().get_or_add_rFonts().attrib.pop(qn('w:' + attr), None)
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.05
    doc.styles['Title'].font.size = Pt(16)
    doc.styles['Title'].paragraph_format.space_after = Pt(14)
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    doc.core_properties.title = title
    doc.core_properties.author = 'Liang Hanghao; Xuan Wentao; Chen Qile; Peng Peng'
    doc.core_properties.last_modified_by = ''
    doc.core_properties.comments = ''
    return doc


def main():
    OUT.mkdir(exist_ok=True)
    cover = new_document('Cover letter')
    cover.add_paragraph('Cover letter', 'Title')
    for block in (OUT / 'cover-letter.txt').read_text(encoding='utf-8-sig').strip().split('\n\n'):
        cover.add_paragraph(block)
    cover.save(OUT / 'cover-letter.docx')

    title = re.search(r'\\title\{([^}]+)\}', (ROOT / 'main.tex').read_text(encoding='utf-8')).group(1)
    highlights = [line.removeprefix('- ').strip() for line in (ROOT / 'highlights.txt').read_text(encoding='utf-8-sig').splitlines() if line.strip()]
    assert 3 <= len(highlights) <= 5 and all(len(line) <= 85 for line in highlights)
    doc = new_document('Highlights')
    doc.add_paragraph('Highlights', 'Title')
    doc.add_paragraph(title).runs[0].bold = True
    for line in highlights:
        doc.add_paragraph(line, 'List Bullet')
    doc.save(OUT / 'highlights.docx')
    print({'outputs': ['cover-letter.docx', 'highlights.docx'], 'highlight_characters': [len(x) for x in highlights]})


if __name__ == '__main__':
    main()
