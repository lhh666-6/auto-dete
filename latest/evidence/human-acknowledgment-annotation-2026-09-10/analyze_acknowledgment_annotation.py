from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Iterable

import numpy as np
import openpyxl


EXPECTED_ROWS = 325
EXPECTED_ENDPOINTS = {"context": 174, "stale": 151}
EXPECTED_SCENARIOS = {"A2": 86, "A3": 88, "A6": 69, "B4": 82}
METRIC_NAMES = ("percent_agreement", "cohen_kappa", "gwet_ac1", "pabak")


def normalize_text(value: str) -> str:
    if not isinstance(value, str):
        raise ValueError("assistant_text must be a string")
    return value.replace("\r\n", "\n").replace("\r", "\n").strip()


def parse_bool(value: Any) -> int:
    if isinstance(value, bool):
        return int(value)
    normalized = str(value).strip().lower()
    if normalized in {"true", "1"}:
        return 1
    if normalized in {"false", "0"}:
        return 0
    raise ValueError(f"not a Boolean label: {value!r}")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def build_exact_text_index(rows: Iterable[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    index: dict[str, dict[str, Any]] = {}
    for row in rows:
        key = normalize_text(row["assistant_text"])
        if key in index:
            raise ValueError("duplicate normalized assistant_text in frozen source audit")
        index[key] = row
    return index


def read_workbook_rows(path: Path) -> list[dict[str, Any]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if "正式标注" not in workbook.sheetnames:
            raise ValueError("workbook is missing the 正式标注 sheet")
        sheet = workbook["正式标注"]
        headers = [sheet.cell(1, column).value for column in range(1, sheet.max_column + 1)]
        required = [
            "case_id",
            "endpoint",
            "assistant_text",
            "label_0_or_1",
            "rationale",
            "uncertainty_note",
        ]
        if headers != required:
            raise ValueError(f"unexpected workbook headers: {headers!r}")
        rows = [
            dict(
                zip(
                    headers,
                    [sheet.cell(row, column).value for column in range(1, sheet.max_column + 1)],
                )
            )
            for row in range(2, sheet.max_row + 1)
        ]
    finally:
        workbook.close()

    if len(rows) != EXPECTED_ROWS:
        raise ValueError(f"expected {EXPECTED_ROWS} workbook rows, found {len(rows)}")
    case_ids = [str(row["case_id"]).strip() for row in rows]
    if len(set(case_ids)) != EXPECTED_ROWS:
        raise ValueError("workbook case_id values are not unique")

    for row, case_id in zip(rows, case_ids):
        row["case_id"] = case_id
        endpoint = str(row["endpoint"]).strip()
        if endpoint not in EXPECTED_ENDPOINTS:
            raise ValueError(f"unexpected endpoint for {case_id}: {endpoint!r}")
        row["endpoint"] = endpoint
        label = row["label_0_or_1"]
        if isinstance(label, bool) or not isinstance(label, int) or label not in {0, 1}:
            raise ValueError(f"non-binary label for {case_id}: {label!r}")
        if row["rationale"] is None or not str(row["rationale"]).strip():
            raise ValueError(f"missing rationale for {case_id}")
        row["assistant_text"] = normalize_text(row["assistant_text"])

    endpoint_counts = Counter(row["endpoint"] for row in rows)
    if dict(endpoint_counts) != EXPECTED_ENDPOINTS:
        raise ValueError(f"unexpected endpoint denominators: {dict(endpoint_counts)!r}")
    return rows


def load_and_join(
    workbook_path: Path,
    source_audit_path: Path,
    repaired_audit_path: Path,
) -> list[dict[str, Any]]:
    workbook_rows = read_workbook_rows(workbook_path)
    source_rows = read_csv(source_audit_path)
    repaired_rows = read_csv(repaired_audit_path)
    if len(source_rows) != EXPECTED_ROWS or len(repaired_rows) != EXPECTED_ROWS:
        raise ValueError(
            f"expected {EXPECTED_ROWS} rows in each audit, found "
            f"{len(source_rows)} and {len(repaired_rows)}"
        )

    source_by_text = build_exact_text_index(source_rows)
    repaired_by_run: dict[str, dict[str, str]] = {}
    for row in repaired_rows:
        run_id = row["run_id"]
        if run_id in repaired_by_run:
            raise ValueError(f"duplicate run_id in repaired audit: {run_id}")
        repaired_by_run[run_id] = row

    joined: list[dict[str, Any]] = []
    seen_runs: set[str] = set()
    for human in workbook_rows:
        source = source_by_text.get(human["assistant_text"])
        if source is None:
            raise ValueError(f"no exact source match for {human['case_id']}")
        run_id = source["run_id"]
        if run_id in seen_runs:
            raise ValueError(f"multiple workbook rows matched run_id {run_id}")
        seen_runs.add(run_id)
        repaired = repaired_by_run.get(run_id)
        if repaired is None:
            raise ValueError(f"no repaired audit row for run_id {run_id}")
        if human["endpoint"] != source["metric"] or source["metric"] != repaired["metric"]:
            raise ValueError(f"endpoint mismatch for {human['case_id']}")
        for field in ("scenario_id", "prompt_variant_id", "repetition", "model_config_id"):
            if str(source[field]) != str(repaired[field]):
                raise ValueError(f"audit mismatch for {run_id}: {field}")

        joined.append(
            {
                "case_id": human["case_id"],
                "endpoint": human["endpoint"],
                "run_id": run_id,
                "scenario_id": source["scenario_id"],
                "prompt_variant_id": source["prompt_variant_id"],
                "repetition": int(source["repetition"]),
                "model_config_id": source["model_config_id"],
                "human_label": int(human["label_0_or_1"]),
                "rationale": str(human["rationale"]),
                "uncertainty_note": ""
                if human["uncertainty_note"] is None
                else str(human["uncertainty_note"]),
                "historical_label": parse_bool(repaired["old_label"]),
                "v1_label": parse_bool(repaired["v1_label"]),
                "v1_rule": repaired["v1_rule"],
                "v2_label": parse_bool(repaired["v2_label"]),
                "v2_rule": repaired["v2_rule"],
                "v2_changed": parse_bool(repaired["changed"]),
            }
        )

    if len(seen_runs) != EXPECTED_ROWS:
        raise ValueError("the joined run_id set is incomplete")
    scenario_counts = Counter(row["scenario_id"] for row in joined)
    if dict(scenario_counts) != EXPECTED_SCENARIOS:
        raise ValueError(f"unexpected scenario denominators: {dict(scenario_counts)!r}")
    return joined


def agreement_metrics(rows: list[dict[str, Any]], label_field: str) -> dict[str, Any]:
    if not rows:
        raise ValueError("agreement metrics require at least one row")
    pairs = [(int(row[label_field]), int(row["human_label"])) for row in rows]
    if any(rule not in {0, 1} or human not in {0, 1} for rule, human in pairs):
        raise ValueError("agreement metrics require binary labels")
    cells = Counter(pairs)
    n = len(pairs)
    agreement_count = cells[(0, 0)] + cells[(1, 1)]
    observed = agreement_count / n
    rule_positive = sum(rule for rule, _ in pairs)
    human_positive = sum(human for _, human in pairs)
    rule_prevalence = rule_positive / n
    human_prevalence = human_positive / n
    expected = rule_prevalence * human_prevalence + (1 - rule_prevalence) * (
        1 - human_prevalence
    )
    kappa = None if math.isclose(expected, 1.0) else (observed - expected) / (1 - expected)
    average_positive = (rule_prevalence + human_prevalence) / 2
    ac1_expected = 2 * average_positive * (1 - average_positive)
    ac1 = None if math.isclose(ac1_expected, 1.0) else (observed - ac1_expected) / (
        1 - ac1_expected
    )
    return {
        "n": n,
        "rule_0_human_0": cells[(0, 0)],
        "rule_0_human_1": cells[(0, 1)],
        "rule_1_human_0": cells[(1, 0)],
        "rule_1_human_1": cells[(1, 1)],
        "agreement_count": agreement_count,
        "percent_agreement": observed,
        "cohen_kappa": kappa,
        "gwet_ac1": ac1,
        "pabak": 2 * observed - 1,
        "human_positive": human_positive,
        "human_positive_rate": human_prevalence,
        "rule_positive": rule_positive,
        "rule_positive_rate": rule_prevalence,
    }


def _resample(
    rows: list[dict[str, Any]], unit: str, generator: np.random.Generator
) -> list[dict[str, Any]]:
    if unit == "case":
        indices = generator.integers(0, len(rows), size=len(rows))
        return [rows[int(index)] for index in indices]
    if unit == "model":
        key = lambda row: row["model_config_id"]
    elif unit == "cell":
        key = lambda row: (row["model_config_id"], row["scenario_id"])
    else:
        raise ValueError(f"unknown bootstrap unit: {unit}")
    clusters: dict[Any, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        clusters[key(row)].append(row)
    names = sorted(clusters, key=str)
    selected = generator.integers(0, len(names), size=len(names))
    return [row for index in selected for row in clusters[names[int(index)]]]


def bootstrap_intervals(
    rows: list[dict[str, Any]],
    label_field: str,
    unit: str,
    replicates: int,
    seed: int,
) -> dict[str, Any]:
    if replicates <= 0:
        raise ValueError("bootstrap replicates must be positive")
    generator = np.random.default_rng(seed)
    values: dict[str, list[float]] = {name: [] for name in METRIC_NAMES}
    undefined: Counter[str] = Counter()
    for _ in range(replicates):
        metrics = agreement_metrics(_resample(rows, unit, generator), label_field)
        for name in METRIC_NAMES:
            value = metrics[name]
            if value is None or not math.isfinite(value):
                undefined[name] += 1
            else:
                values[name].append(float(value))
    intervals: dict[str, Any] = {}
    for name in METRIC_NAMES:
        samples = values[name]
        intervals[name] = {
            "lower_95": None if not samples else float(np.quantile(samples, 0.025)),
            "upper_95": None if not samples else float(np.quantile(samples, 0.975)),
            "valid_replicates": len(samples),
            "undefined_replicates": int(undefined[name]),
        }
    return {
        "unit": unit,
        "replicates": replicates,
        "seed": seed,
        "intervals": intervals,
    }


def strata(rows: list[dict[str, Any]]) -> list[tuple[str, str, list[dict[str, Any]]]]:
    result = [("overall", "all", rows)]
    for endpoint in ("context", "stale"):
        result.append(("endpoint", endpoint, [row for row in rows if row["endpoint"] == endpoint]))
    for scenario in ("A2", "A3", "A6", "B4"):
        result.append(
            ("scenario", scenario, [row for row in rows if row["scenario_id"] == scenario])
        )
    for model in ("D1", "G1", "G2"):
        result.append(
            ("model_config", model, [row for row in rows if row["model_config_id"] == model])
        )
    return result


def write_csv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_latex_table(
    path: Path,
    summary_rows: list[dict[str, Any]],
    bootstrap_rows: list[dict[str, Any]],
) -> None:
    scenario_rows = [
        row
        for row in summary_rows
        if row["comparison"] == "v2.1 rule" and row["stratum_type"] in {"overall", "scenario"}
    ]
    labels = {"all": "Overall", "A2": "A2", "A3": "A3 (exploratory)", "A6": "A6", "B4": "B4"}
    lines = [
        "% Generated by analyze_acknowledgment_annotation.py; do not edit by hand.",
        "\\begin{table}[t]",
        "\\centering",
        "\\caption{Blinded author annotation versus the repaired v2.1 textual rule. Agreement is human--rule agreement, not inter-human reliability; $\\kappa$ is Cohen's $\\kappa$.}",
        "\\label{tab:acknowledgment-human-rule}",
        "\\begingroup",
        "\\footnotesize",
        "\\setlength{\\tabcolsep}{4pt}",
        "\\begin{tabular}{@{}lrrrrrrr@{}}",
        "\\toprule",
        "Stratum & $n$ & H+ & R+ & Agree & $\\kappa$ & AC1 & PABAK \\\\",
        "\\midrule",
    ]
    for row in scenario_rows:
        kappa = "--" if row["cohen_kappa"] is None else f"{row['cohen_kappa']:.3f}"
        ac1 = "--" if row["gwet_ac1"] is None else f"{row['gwet_ac1']:.3f}"
        lines.append(
            f"{labels[row['stratum']]} & {row['n']} & {row['human_positive']} & {row['rule_positive']} "
            f"& {100 * row['percent_agreement']:.1f}\\% & {kappa} & {ac1} & {row['pabak']:.3f} \\\\")
    lines.extend(["\\bottomrule", "\\end{tabular}", "\\endgroup", "\\end{table}", ""])

    lines.extend(
        [
            "\\begin{table}[t]",
            "\\centering",
            "\\caption{Confusion counts for the blinded author annotation and repaired v2.1 rule. Rule 0/H 0 denotes rule label 0 and human label 0; the remaining R/H columns follow the same order. No disagreements were adjudicated.}",
            "\\label{tab:acknowledgment-confusion}",
            "\\begingroup",
            "\\footnotesize",
            "\\setlength{\\tabcolsep}{4pt}",
        "\\begin{tabular}{@{}lrrrrrr@{}}",
        "\\toprule",
        "Stratum & $n$ & Agree & R0/H0 & R0/H1 & R1/H0 & R1/H1 \\\\",
        "\\midrule",
        ]
    )
    for row in scenario_rows:
        lines.append(
            f"{labels[row['stratum']]} & {row['n']} & {100 * row['percent_agreement']:.1f}\\% "
            f"& {row['rule_0_human_0']} & {row['rule_0_human_1']} "
            f"& {row['rule_1_human_0']} & {row['rule_1_human_1']} \\\\")
    lines.extend(["\\bottomrule", "\\end{tabular}", "\\endgroup", "\\end{table}", ""])

    lookup = {
        (row["unit"], row["metric"]): row
        for row in bootstrap_rows
        if row["stratum_type"] == "overall" and row["stratum"] == "all"
    }
    unit_labels = {"case": "Case resampling", "model": "Model resampling", "cell": "Cell resampling"}
    lines.extend(
        [
            "\\begin{table}[t]",
            "\\centering",
            "\\caption{Overall human--rule bootstrap sensitivity. Intervals are percentile 95\\% intervals from 5,000 replicates. Model resampling has three clusters; cell resampling has twelve model$\\times$scenario clusters.}",
            "\\label{tab:acknowledgment-bootstrap}",
            "\\small",
            "\\begin{tabular}{lrrr}",
            "\\toprule",
            "Unit & Agreement 95\\% CI & $\\kappa$ 95\\% CI & Valid $\\kappa$ replicates \\\\",
            "\\midrule",
        ]
    )
    for unit in ("case", "model", "cell"):
        agreement = lookup[(unit, "percent_agreement")]
        kappa = lookup[(unit, "cohen_kappa")]
        lines.append(
            f"{unit_labels[unit]} & [{100 * float(agreement['lower_95']):.1f}\\%, "
            f"{100 * float(agreement['upper_95']):.1f}\\%] & "
            f"[{float(kappa['lower_95']):.3f}, {float(kappa['upper_95']):.3f}] & "
            f"{kappa['valid_replicates']} \\\\")
    lines.extend(["\\bottomrule", "\\end{tabular}", "\\end{table}", ""])
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")


def run_analysis(
    workbook_path: Path,
    source_audit_path: Path,
    repaired_audit_path: Path,
    output_dir: Path,
    *,
    bootstrap_replicates: int = 5000,
    bootstrap_seed: int = 20260910,
    latex_table_path: Path | None = None,
) -> dict[str, Any]:
    output_dir.mkdir(parents=True, exist_ok=True)
    rows = load_and_join(workbook_path, source_audit_path, repaired_audit_path)
    normalized_fields = [
        "case_id",
        "endpoint",
        "run_id",
        "scenario_id",
        "prompt_variant_id",
        "repetition",
        "model_config_id",
        "human_label",
        "rationale",
        "uncertainty_note",
        "historical_label",
        "v1_label",
        "v1_rule",
        "v2_label",
        "v2_rule",
        "v2_changed",
    ]
    write_csv(output_dir / "normalized-labels.csv", rows, normalized_fields)

    summary_rows: list[dict[str, Any]] = []
    comparisons = (("v1 rule", "v1_label"), ("v2.1 rule", "v2_label"))
    grouped = strata(rows)
    for comparison, field in comparisons:
        for stratum_type, stratum_name, subset in grouped:
            summary_rows.append(
                {
                    "comparison": comparison,
                    "stratum_type": stratum_type,
                    "stratum": stratum_name,
                    **agreement_metrics(subset, field),
                }
            )
    summary_fields = list(summary_rows[0])
    write_csv(output_dir / "agreement-summary.csv", summary_rows, summary_fields)

    disagreements = [row for row in rows if row["human_label"] != row["v2_label"]]
    write_csv(output_dir / "disagreements-v2.csv", disagreements, normalized_fields)

    bootstrap_rows: list[dict[str, Any]] = []
    for stratum_index, (stratum_type, stratum_name, subset) in enumerate(grouped):
        for unit_index, unit in enumerate(("case", "model", "cell")):
            result = bootstrap_intervals(
                subset,
                "v2_label",
                unit,
                bootstrap_replicates,
                bootstrap_seed + stratum_index * 10 + unit_index,
            )
            for metric, interval in result["intervals"].items():
                bootstrap_rows.append(
                    {
                        "stratum_type": stratum_type,
                        "stratum": stratum_name,
                        "unit": unit,
                        "metric": metric,
                        "replicates": result["replicates"],
                        "seed": result["seed"],
                        **interval,
                    }
                )
    write_csv(output_dir / "bootstrap-intervals.csv", bootstrap_rows, list(bootstrap_rows[0]))

    v2_by_stratum = {
        f"{row['stratum_type']}:{row['stratum']}": {
            key: value
            for key, value in row.items()
            if key not in {"comparison", "stratum_type", "stratum"}
        }
        for row in summary_rows
        if row["comparison"] == "v2.1 rule"
    }
    report = {
        "schema": "auto-decte.human-acknowledgment-annotation.v1",
        "annotation": {
            "annotator": "second author Xuan Wentao",
            "design": "author-involved blinded manual annotation",
            "third_party_independent": False,
            "human_human_reliability": False,
            "adjudicated": False,
            "rows": len(rows),
            "uncertainty_notes": sum(bool(row["uncertainty_note"].strip()) for row in rows),
        },
        "inputs": {
            "completed_workbook": {
                "file": workbook_path.name,
                "sha256": sha256_file(workbook_path),
            },
            "source_semantic_audit": {
                "file": source_audit_path.name,
                "sha256": sha256_file(source_audit_path),
            },
            "repaired_v2_audit": {
                "file": repaired_audit_path.name,
                "sha256": sha256_file(repaired_audit_path),
            },
        },
        "join": {
            "method": "exact assistant_text after line-ending normalization and trim",
            "matched": len(rows),
            "missing": 0,
            "ambiguous": 0,
        },
        "primary_comparison": {
            "label": "manual label versus repaired v2.1 rule label",
            "overall": v2_by_stratum["overall:all"],
            "by_endpoint": {
                name: v2_by_stratum[f"endpoint:{name}"] for name in ("context", "stale")
            },
            "by_scenario": {
                name: v2_by_stratum[f"scenario:{name}"] for name in ("A2", "A3", "A6", "B4")
            },
            "by_model_config": {
                name: v2_by_stratum[f"model_config:{name}"] for name in ("D1", "G1", "G2")
            },
            "disagreements": len(disagreements),
        },
        "bootstrap": {
            "replicates": bootstrap_replicates,
            "base_seed": bootstrap_seed,
            "units": ["case", "model", "model-by-scenario cell"],
            "file": "bootstrap-intervals.csv",
        },
        "interpretation": {
            "overall_is_descriptive": True,
            "a3_is_exploratory": True,
            "reason": "A3 did not specify a unique intended target field.",
            "rule_labels_overwritten": False,
        },
    }
    (output_dir / "analysis-report.json").write_text(
        json.dumps(report, indent=2, ensure_ascii=False, sort_keys=True) + "\n",
        encoding="utf-8",
    )

    scenario_lines = []
    for name in ("A2", "A3", "A6", "B4"):
        metrics = report["primary_comparison"]["by_scenario"][name]
        scenario_lines.append(
            f"| {name}{' (exploratory)' if name == 'A3' else ''} | {metrics['n']} | "
            f"{metrics['human_positive']} | {metrics['rule_positive']} | "
            f"{metrics['agreement_count']}/{metrics['n']} ({100 * metrics['percent_agreement']:.2f}%) | "
            f"{metrics['cohen_kappa']:.3f} |"
        )
    overall = report["primary_comparison"]["overall"]
    markdown = f"""# Blinded acknowledgment annotation analysis

The second author, Xuan Wentao, independently assigned the manual labels while blinded to previous labels and configuration/run/scenario identifiers. Because the annotator is an author, this is author-involved blinded annotation rather than independent third-party annotation. The comparison below is human--rule agreement, not inter-human reliability or adjudicated truth.

All {len(rows)} workbook rows matched exactly one frozen audit row. The manual annotation contains {sum(row['human_label'] for row in rows)} positive labels; the repaired v2.1 rule contains {sum(row['v2_label'] for row in rows)} rule hits.

## Primary comparison

Overall agreement is {overall['agreement_count']}/{overall['n']} ({100 * overall['percent_agreement']:.2f}%); Cohen's kappa is {overall['cohen_kappa']:.3f}, Gwet's AC1 is {overall['gwet_ac1']:.3f}, and PABAK is {overall['pabak']:.3f}.

| Scenario | n | Human positive | Rule positive | Agreement | Kappa |
|---|---:|---:|---:|---:|---:|
{chr(10).join(scenario_lines)}

The A3 result is not interpreted as annotator error or model failure. Its frozen prompt did not identify a unique target field, so its weaker agreement exposes construct and rubric ambiguity. The A2, A3, A6, and B4 results remain separate.

## Retained outputs

The archive retains normalized joined labels, both historical and repaired comparisons, all {len(disagreements)} unadjudicated v2.1 disagreements, and 5,000-replicate case/model/cell bootstrap intervals. The original rule labels are not overwritten.
"""
    (output_dir / "analysis-report.md").write_text(markdown, encoding="utf-8")
    if latex_table_path is not None:
        write_latex_table(latex_table_path, summary_rows, bootstrap_rows)
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--source-audit", type=Path, required=True)
    parser.add_argument("--repaired-audit", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--latex-table", type=Path)
    parser.add_argument("--bootstrap-replicates", type=int, default=5000)
    parser.add_argument("--bootstrap-seed", type=int, default=20260910)
    args = parser.parse_args()
    report = run_analysis(
        args.workbook,
        args.source_audit,
        args.repaired_audit,
        args.output_dir,
        bootstrap_replicates=args.bootstrap_replicates,
        bootstrap_seed=args.bootstrap_seed,
        latex_table_path=args.latex_table,
    )
    overall = report["primary_comparison"]["overall"]
    print(
        f"matched={report['join']['matched']} agreement="
        f"{overall['agreement_count']}/{overall['n']} "
        f"kappa={overall['cohen_kappa']:.6f}"
    )


if __name__ == "__main__":
    main()
