"""Traceable four-sheet XLSX writer."""

from collections.abc import Iterable
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from app.application.query_forms import SearchResult


class XlsxExporter:
    def write(
        self,
        destination: Path,
        batch_id: str,
        export_type: str,
        results: Iterable[SearchResult],
        filters: dict[str, Any],
    ) -> None:
        rows = list(results)
        workbook = Workbook()
        official = workbook.active
        assert official is not None
        official.title = "正式数据"
        field_names = sorted({key for result in rows for key in result.current_record.values})
        official.append(["export_batch_id", "form_id", "record_version", *field_names])
        for result in rows:
            official.append(
                [
                    batch_id,
                    result.form.form_id,
                    result.current_record.version,
                    *(result.current_record.values.get(name) for name in field_names),
                ]
            )

        review = workbook.create_sheet("异常与复核")
        review.append(["form_id", "record_version", "status", "change_reason", "confirmed_by"])
        for result in rows:
            record = result.current_record
            review.append(
                [
                    result.form.form_id,
                    record.version,
                    record.status.value,
                    record.change_reason,
                    record.confirmed_by,
                ]
            )

        summary = workbook.create_sheet("汇总")
        summary.append(["export_type", "record_count", "total_quantity", "qualified_quantity"])
        summary.append(
            [
                export_type,
                len(rows),
                sum(int(row.current_record.values.get("total_quantity", 0)) for row in rows),
                sum(int(row.current_record.values.get("qualified_quantity", 0)) for row in rows),
            ]
        )

        information = workbook.create_sheet("导出说明")
        information.append(["export_batch_id", batch_id])
        information.append(["export_type", export_type])
        information.append(["filters", str(filters)])
        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(destination)
